import os
import shutil
from decimal import Decimal
from typing import List, Optional
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse, FileResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session
import io

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads", "contabilita")
os.makedirs(UPLOAD_DIR, exist_ok=True)

ALLOWED_MIME = {
    "image/jpeg", "image/png", "image/gif", "image/webp",
    "application/pdf",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
MAX_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

from app.database import get_db
from app import models, schemas
from app.auth import get_current_user

router = APIRouter(prefix="/contabilita", tags=["contabilita"])


# ─── Categorie ───────────────────────────────────────────────────────────────

@router.get("/categorie/", response_model=List[schemas.CategoriaContabileResponse])
def list_categorie(
    tipo: Optional[str] = None,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    q = db.query(models.CategoriaContabile)
    if tipo:
        q = q.filter(models.CategoriaContabile.tipo == tipo)
    return q.order_by(models.CategoriaContabile.tipo, models.CategoriaContabile.nome).all()


@router.post("/categorie/", response_model=schemas.CategoriaContabileResponse, status_code=201)
def create_categoria(
    data: schemas.CategoriaContabileCreate,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = models.CategoriaContabile(**data.model_dump())
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@router.put("/categorie/{id}", response_model=schemas.CategoriaContabileResponse)
def update_categoria(
    id: int,
    data: schemas.CategoriaContabileUpdate,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.CategoriaContabile).filter(models.CategoriaContabile.id == id).first()
    if not obj:
        raise HTTPException(404, "Categoria non trovata")
    for k, v in data.model_dump(exclude_unset=True).items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return obj


@router.delete("/categorie/{id}", status_code=204)
def delete_categoria(
    id: int,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.CategoriaContabile).filter(models.CategoriaContabile.id == id).first()
    if not obj:
        raise HTTPException(404, "Categoria non trovata")
    db.delete(obj)
    db.commit()


# ─── Fornitori ────────────────────────────────────────────────────────────────

@router.get("/fornitori/", response_model=List[schemas.FornitoreResponse])
def list_fornitori(
    search: Optional[str] = None,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    q = db.query(models.Fornitore)
    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            models.Fornitore.nome.ilike(like),
            models.Fornitore.partita_iva.ilike(like),
        ))
    return q.order_by(models.Fornitore.nome).all()


@router.post("/fornitori/", response_model=schemas.FornitoreResponse, status_code=201)
def create_fornitore(
    data: schemas.FornitoreCreate,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = models.Fornitore(**data.model_dump())
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@router.put("/fornitori/{id}", response_model=schemas.FornitoreResponse)
def update_fornitore(
    id: int,
    data: schemas.FornitoreUpdate,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.Fornitore).filter(models.Fornitore.id == id).first()
    if not obj:
        raise HTTPException(404, "Fornitore non trovato")
    for k, v in data.model_dump(exclude_unset=True).items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return obj


@router.delete("/fornitori/{id}", status_code=204)
def delete_fornitore(
    id: int,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.Fornitore).filter(models.Fornitore.id == id).first()
    if not obj:
        raise HTTPException(404, "Fornitore non trovato")
    db.delete(obj)
    db.commit()


# ─── Movimenti ────────────────────────────────────────────────────────────────

@router.get("/movimenti/", response_model=List[schemas.MovimentoContabileResponse])
def list_movimenti(
    search: Optional[str] = None,
    tipo: Optional[str] = None,
    categoria_id: Optional[int] = None,
    fornitore_id: Optional[int] = None,
    dal: Optional[date] = None,
    al: Optional[date] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    q = db.query(models.MovimentoContabile)
    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            models.MovimentoContabile.descrizione.ilike(like),
            models.MovimentoContabile.numero_documento.ilike(like),
        ))
    if tipo:
        q = q.filter(models.MovimentoContabile.tipo == tipo)
    if categoria_id:
        q = q.filter(models.MovimentoContabile.categoria_id == categoria_id)
    if fornitore_id:
        q = q.filter(models.MovimentoContabile.fornitore_id == fornitore_id)
    if dal:
        q = q.filter(models.MovimentoContabile.data >= dal)
    if al:
        q = q.filter(models.MovimentoContabile.data <= al)

    movimenti = q.order_by(models.MovimentoContabile.data.desc()).offset(skip).limit(limit).all()

    # Hydrate categoria_nome and fornitore_nome
    result = []
    for m in movimenti:
        d = schemas.MovimentoContabileResponse.model_validate(m)
        if m.categoria_id:
            cat = db.query(models.CategoriaContabile).filter(models.CategoriaContabile.id == m.categoria_id).first()
            d.categoria_nome = cat.nome if cat else None
        if m.fornitore_id:
            forn = db.query(models.Fornitore).filter(models.Fornitore.id == m.fornitore_id).first()
            d.fornitore_nome = forn.nome if forn else None
        result.append(d)
    return result


@router.post("/movimenti/", response_model=schemas.MovimentoContabileResponse, status_code=201)
def create_movimento(
    data: schemas.MovimentoContabileCreate,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = models.MovimentoContabile(**data.model_dump())
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return schemas.MovimentoContabileResponse.model_validate(obj)


@router.get("/movimenti/{id}", response_model=schemas.MovimentoContabileResponse)
def get_movimento(
    id: int,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    m = db.query(models.MovimentoContabile).filter(models.MovimentoContabile.id == id).first()
    if not m:
        raise HTTPException(404, "Movimento non trovato")
    d = schemas.MovimentoContabileResponse.model_validate(m)
    if m.categoria_id:
        cat = db.query(models.CategoriaContabile).filter(models.CategoriaContabile.id == m.categoria_id).first()
        d.categoria_nome = cat.nome if cat else None
    if m.fornitore_id:
        forn = db.query(models.Fornitore).filter(models.Fornitore.id == m.fornitore_id).first()
        d.fornitore_nome = forn.nome if forn else None
    return d


@router.put("/movimenti/{id}", response_model=schemas.MovimentoContabileResponse)
def update_movimento(
    id: int,
    data: schemas.MovimentoContabileUpdate,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.MovimentoContabile).filter(models.MovimentoContabile.id == id).first()
    if not obj:
        raise HTTPException(404, "Movimento non trovato")
    for k, v in data.model_dump(exclude_unset=True).items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return schemas.MovimentoContabileResponse.model_validate(obj)


@router.delete("/movimenti/{id}", status_code=204)
def delete_movimento(
    id: int,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.MovimentoContabile).filter(models.MovimentoContabile.id == id).first()
    if not obj:
        raise HTTPException(404, "Movimento non trovato")
    if obj.allegato_path and os.path.exists(obj.allegato_path):
        os.remove(obj.allegato_path)
    db.delete(obj)
    db.commit()


# ─── Allegati ─────────────────────────────────────────────────────────────────

@router.post("/movimenti/{id}/allegato", response_model=schemas.MovimentoContabileResponse)
async def upload_allegato(
    id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.MovimentoContabile).filter(models.MovimentoContabile.id == id).first()
    if not obj:
        raise HTTPException(404, "Movimento non trovato")

    if file.content_type not in ALLOWED_MIME:
        raise HTTPException(400, f"Tipo file non supportato: {file.content_type}")

    content = await file.read()
    if len(content) > MAX_SIZE_BYTES:
        raise HTTPException(400, "File troppo grande (max 10 MB)")

    # Remove old attachment if present
    if obj.allegato_path and os.path.exists(obj.allegato_path):
        os.remove(obj.allegato_path)

    # Sanitise filename and save
    safe_name = os.path.basename(file.filename or "allegato")
    dest = os.path.join(UPLOAD_DIR, f"{id}_{safe_name}")
    with open(dest, "wb") as f:
        f.write(content)

    obj.allegato_path = dest
    obj.allegato_nome = file.filename
    db.commit()
    db.refresh(obj)
    return schemas.MovimentoContabileResponse.model_validate(obj)


@router.get("/movimenti/{id}/allegato")
def download_allegato(
    id: int,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.MovimentoContabile).filter(models.MovimentoContabile.id == id).first()
    if not obj or not obj.allegato_path:
        raise HTTPException(404, "Allegato non trovato")
    if not os.path.exists(obj.allegato_path):
        raise HTTPException(404, "File non trovato sul disco")
    return FileResponse(
        obj.allegato_path,
        filename=obj.allegato_nome or os.path.basename(obj.allegato_path),
    )


@router.delete("/movimenti/{id}/allegato", status_code=204)
def delete_allegato(
    id: int,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    obj = db.query(models.MovimentoContabile).filter(models.MovimentoContabile.id == id).first()
    if not obj:
        raise HTTPException(404, "Movimento non trovato")
    if obj.allegato_path and os.path.exists(obj.allegato_path):
        os.remove(obj.allegato_path)
    obj.allegato_path = None
    obj.allegato_nome = None
    db.commit()


# ─── Riepilogo ────────────────────────────────────────────────────────────────

@router.get("/riepilogo/")
def riepilogo(
    dal: Optional[date] = None,
    al: Optional[date] = None,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    q = db.query(models.MovimentoContabile)
    if dal:
        q = q.filter(models.MovimentoContabile.data >= dal)
    if al:
        q = q.filter(models.MovimentoContabile.data <= al)

    movimenti = q.all()

    tot_entrate = sum(float(m.importo) for m in movimenti if m.tipo == 'entrata')
    tot_uscite = sum(float(m.importo) for m in movimenti if m.tipo == 'uscita')

    # Per categoria
    cat_totali: dict = {}
    for m in movimenti:
        key = m.categoria_id
        if key not in cat_totali:
            cat = db.query(models.CategoriaContabile).filter(models.CategoriaContabile.id == key).first() if key else None
            cat_totali[key] = {
                "categoria_id": key,
                "categoria": cat.nome if cat else "Senza categoria",
                "tipo": m.tipo,
                "colore": cat.colore if cat else "#9ca3af",
                "totale": 0.0,
            }
        cat_totali[key]["totale"] += float(m.importo)

    return {
        "totale_entrate": round(tot_entrate, 2),
        "totale_uscite": round(tot_uscite, 2),
        "saldo": round(tot_entrate - tot_uscite, 2),
        "per_categoria": sorted(cat_totali.values(), key=lambda x: -x["totale"]),
    }


# ─── Export Excel ─────────────────────────────────────────────────────────────

@router.get("/export/excel")
def export_excel(
    dal: Optional[date] = None,
    al: Optional[date] = None,
    db: Session = Depends(get_db),
    _: models.Utente = Depends(get_current_user),
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    q = db.query(models.MovimentoContabile)
    if dal:
        q = q.filter(models.MovimentoContabile.data >= dal)
    if al:
        q = q.filter(models.MovimentoContabile.data <= al)
    movimenti = q.order_by(models.MovimentoContabile.data.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimenti"

    # Header
    headers = ["Data", "Tipo", "Descrizione", "Importo (€)", "Categoria", "Fornitore", "N° Documento", "Metodo Pagamento", "Note"]
    header_fill = PatternFill(start_color="1e3a5f", end_color="1e3a5f", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Rows
    for row_idx, m in enumerate(movimenti, 2):
        cat = db.query(models.CategoriaContabile).filter(models.CategoriaContabile.id == m.categoria_id).first() if m.categoria_id else None
        forn = db.query(models.Fornitore).filter(models.Fornitore.id == m.fornitore_id).first() if m.fornitore_id else None

        ws.cell(row=row_idx, column=1, value=m.data.isoformat() if m.data else "")
        ws.cell(row=row_idx, column=2, value="Entrata" if m.tipo == "entrata" else "Uscita")
        ws.cell(row=row_idx, column=3, value=m.descrizione)
        cell_importo = ws.cell(row=row_idx, column=4, value=float(m.importo))
        cell_importo.number_format = '#,##0.00'
        if m.tipo == "entrata":
            cell_importo.font = Font(color="166534")  # green
        else:
            cell_importo.font = Font(color="991b1b")  # red
        ws.cell(row=row_idx, column=5, value=cat.nome if cat else "")
        ws.cell(row=row_idx, column=6, value=forn.nome if forn else "")
        ws.cell(row=row_idx, column=7, value=m.numero_documento or "")
        ws.cell(row=row_idx, column=8, value=m.metodo_pagamento or "")
        ws.cell(row=row_idx, column=9, value=m.note or "")

    # Auto column widths
    for col in range(1, len(headers) + 1):
        max_len = max(
            (len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 4, 50)

    # Totals row
    tot_row = ws.max_row + 2
    ws.cell(row=tot_row, column=3, value="TOTALE ENTRATE:").font = Font(bold=True)
    ws.cell(row=tot_row, column=4, value=sum(float(m.importo) for m in movimenti if m.tipo == "entrata")).number_format = '#,##0.00'
    ws.cell(row=tot_row + 1, column=3, value="TOTALE USCITE:").font = Font(bold=True)
    ws.cell(row=tot_row + 1, column=4, value=sum(float(m.importo) for m in movimenti if m.tipo == "uscita")).number_format = '#,##0.00'
    ws.cell(row=tot_row + 2, column=3, value="SALDO:").font = Font(bold=True)
    saldo = sum(float(m.importo) for m in movimenti if m.tipo == "entrata") - sum(float(m.importo) for m in movimenti if m.tipo == "uscita")
    ws.cell(row=tot_row + 2, column=4, value=saldo).number_format = '#,##0.00'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"contabilita_{dal or 'inizio'}_{al or 'oggi'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
